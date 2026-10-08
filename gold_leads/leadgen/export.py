"""Local copies of the lead list: Excel (same tabs as the Google Sheet) and CSV."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .table import ALL_TAB, SYNC_TAB


def read_xlsx(path: Path) -> dict[str, list[list]]:
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True)
    tabs = {ws.title: [[("" if v is None else v) for v in row] for row in ws.iter_rows(values_only=True)] for ws in wb}
    wb.close()
    return tabs


def write_xlsx(tabs: dict[str, list[list]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="FBE59A")
    for title, rows in tabs.items():
        ws = wb.create_sheet(title[:31])
        ws.sheet_view.rightToLeft = True
        if title == SYNC_TAB:
            ws.sheet_state = "hidden"
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, longest + 2), 60)
    wb.save(path)


def write_csv(tabs: dict[str, list[list]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # BOM so Excel shows Persian correctly
        csv.writer(f).writerows(tabs.get(ALL_TAB, []))
